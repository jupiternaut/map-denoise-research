"""Build the report from frozen results, without executing experiment code."""
from pathlib import Path
from collections import defaultdict
import csv
import hashlib
import json
import re
import socket

import numpy as np
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

HERE = Path(__file__).resolve().parent
SOURCE = HERE.parent / 'e0-diagnostics-20260915T072044Z'
OLD = HERE.parent / 'e0'
assert socket.gethostname() == 'liekkas'


def digest(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def source_snapshot():
    return {str(p): digest(p) for root in (SOURCE, OLD)
            for p in sorted(root.rglob('*')) if p.is_file()}


def dump(name, value):
    (HERE / name).write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n')


def recheck():
    snapshot = source_snapshot()
    manifest = HERE / 'EVIDENCE_SNAPSHOT.json'
    if manifest.exists():
        assert json.loads(manifest.read_text())['files'] == snapshot, 'Frozen source changed'
    else:
        dump(manifest.name, dict(host=socket.gethostname(), files=snapshot))
    detector = json.loads((SOURCE / 'results/detector/summary.json').read_text())
    rows = [r for r in detector['overall'] if r['selector'] == 'bh']
    assert len(rows) == 18 and sum(r['n_rejected'] == 0 for r in rows) == 17
    nonzero = [r for r in rows if r['n_rejected']]
    assert len(nonzero) == 1 and nonzero[0]['n_rejected'] == 76
    projection = json.loads((SOURCE / 'results/projection/RESULTS.json').read_text())['runs']
    boundary = json.loads((SOURCE / 'results/boundary/RESULTS.json').read_text())['runs']
    bounds = {r['key']: r for r in boundary}
    table = defaultdict(list)
    max_delta = 0.0
    beat_argmin = beat_identity = linear_identity = 0
    for r in projection:
        metrics = r['diagnostics']['eligible_no_bh']['metrics']
        vals = [metrics[k]['ALL']['e_after_mean'] for k in ('identity','projection','argmin')]
        beat_argmin += vals[1] < vals[2]
        beat_identity += vals[1] < vals[0]
        b = bounds[r['key']]
        linear_identity += b['linear']['e_after_mean'] < vals[0]
        table[(r['fbig'],r['calibration'],r['alpha'])].append(
            [vals[0],vals[1],b['linear']['e_after_mean'],vals[2]])
        with np.load(SOURCE / 'results/projection' / (r['key']+'.npz')) as z:
            for arm in ('identity','projection','argmin'):
                out = np.zeros_like(z['evaluation_truth']) if arm == 'identity' else z[arm]
                errors = out-z['evaluation_truth']
                measured = [np.mean(np.abs(errors)), np.sqrt(np.mean(errors**2))]
                saved = [metrics[arm]['ALL']['e_after_mean'],metrics[arm]['ALL']['e_after_rmse']]
                max_delta = max(max_delta, max(abs(x-y) for x,y in zip(measured,saved)))
    assert (beat_argmin,beat_identity,linear_identity) == (24,3,3)
    print('Report-runtime metric recheck maximum delta:',max_delta)
    assert max_delta < 1e-12
    grid_harm = sum(r['grid_covered']['n_dmg'] for r in boundary)
    linear_harm = sum(r['linear_covered']['n_dmg'] for r in boundary)
    assert (grid_harm,linear_harm) == (348,0)
    guards = json.loads((SOURCE / 'results/guards/metrics.json').read_text())
    guard_summary = {}
    for scenario in sorted({r['scenario'] for r in guards}):
        items = [r['arms']['initial_gap_guard'] for r in guards if r['scenario']==scenario]
        guard_summary[scenario] = {k:sum(r[k] for r in items) for k in (
            'proposed_movement_count','alarm_count','beneficial_proposals_retained_count',
            'beneficial_proposals_blocked_count','harmful_proposals_blocked_count')}
    assert guard_summary['wrong_layer_merge']['harmful_proposals_blocked_count'] == 3600
    assert guard_summary['valid_ghost_merge']['beneficial_proposals_blocked_count'] == 3600
    assert sum(r['beneficial_proposals_retained_count'] for r in guard_summary.values()) == 10800
    table_rows = [dict(f_big=k[0], calibration=k[1], alpha=k[2], seeds=len(v),
                       identity=float(np.mean(v,axis=0)[0]), grid=float(np.mean(v,axis=0)[1]),
                       linear=float(np.mean(v,axis=0)[2]), argmin=float(np.mean(v,axis=0)[3]))
                  for k,v in sorted(table.items())]
    md = (HERE/'CPR2_E0D_TECHNICAL_REPORT.md').read_text()
    for row in table_rows:
        for k in ('identity','grid','linear','argmin'):
            assert f"{row[k]:.5f}" in md
    dump('REPORT_RECHECK.json', dict(scope='Recalculation of existing outputs only; no new experiment',
        bh_cells=len(rows), bh_zero_cells=17, nonzero_cell=nonzero,
        projection_configs=len(projection),grid_better_than_argmin=beat_argmin,
        grid_better_than_identity=beat_identity,linear_better_than_identity=linear_identity,
        projection_metric_max_delta=max_delta,projection_table_mm=table_rows,
        covered_grid_harm_events=grid_harm,covered_linear_harm_events=linear_harm,
        max_endpoint_change_mm=max(r['max_movement_difference'] for r in boundary),
        guard_cases=len(guards),guard_summary=guard_summary,
        frozen_files_checked=len(snapshot),frozen_files_unchanged=(snapshot==source_snapshot()),
        old_e0_files=sum(str(p).startswith(str(OLD)+'/') for p in snapshot),
        protocol_sha256=digest(SOURCE/'PROTOCOL.md')))


def node(tag, children=(), **attrs):
    e = OxmlElement('m:'+tag)
    for k,v in attrs.items(): e.set(qn('m:'+k),v)
    for child in children: e.append(child)
    return e


def mr(text):
    r=node('r'); t=node('t'); t.text=text; r.append(t); return r


def sub(base, suffix):
    return node('sSub',[node('e',[base]),node('sub',[mr(suffix)])])


def sup(base, power):
    return node('sSup',[node('e',[base]),node('sup',[mr(power)])])


def both(base, suffix, power):
    return node('sSubSup',[node('e',[base]),node('sub',[mr(suffix)]),node('sup',[mr(power)])])


def hat(base):
    return node('acc',[node('accPr',[node('chr',val='̂')]),node('e',[base])])


def equation(which):
    if which == 0:
        return [sub(mr('s'),'i'),mr('(t) = '),sub(mr('c'),'i'),mr('(t) − '),
                node('limLow',[node('e',[mr('min')]),node('lim',[mr('u')])]),
                sub(mr('c'),'i'),mr('(u),    '),sub(mr('C'),'i'),mr(' = {t : '),
                sub(mr('s'),'i'),mr('(t) ≤ '),sub(mr('q'),'i'),mr('}.')]
    if which == 1:
        summand = [mr('1{'),sub(mr('S'),'ij'),mr(' ≥ '),sub(mr('s'),'i'),mr('(0)}')]
        summation=node('nary',[node('naryPr',[node('chr',val='∑'),node('limLoc',val='subSup')]),
                              node('sub',[mr('j = 1')]),node('sup',[sub(mr('n'),'i')]),node('e',summand)])
        return [sub(mr('p'),'i'),mr(' = '),node('f',[
            node('num',[mr('1 + '),summation]),node('den',[sub(mr('n'),'i'),mr(' + 1')])])]
    if which == 2:
        return [mr('∃ k ∈ {1, …, N} :    '),sub(mr('p'),'(k)'),mr(' ≤ '),
                sub(mr('q'),'BH'),mr(' k / N.')]
    return [sub(hat(mr('t')),'i'),mr(' = '),sub(mr('P'),'Cᵢ'),mr('(0),    '),
            sup(node('box',[node('e',[mr('|'),sub(hat(mr('t')),'i'),mr(' − '),both(mr('t'),'i','*'),mr('|')])]),'2'),
            mr(' ≤ '),sup(node('box',[node('e',[mr('|'),both(mr('t'),'i','*'),mr('|')])]),'2'),
            mr(' − '),sup(node('box',[node('e',[mr('|'),sub(hat(mr('t')),'i'),mr('|')])]),'2'),mr('.')]


def set_font(style, size, east='Noto Serif CJK SC', latin='Liberation Serif'):
    style.font.name=latin; style.font.size=Pt(size); style.font.color.rgb=RGBColor(0,0,0)
    rp=style.element.get_or_add_rPr()
    fonts=rp.find(qn('w:rFonts'))
    if fonts is None:
        fonts=OxmlElement('w:rFonts'); rp.insert(0,fonts)
    for key, val in [('ascii',latin),('hAnsi',latin),('eastAsia',east)]: fonts.set(qn('w:'+key),val)


def inline(p,text):
    # Only formatting supported by the report's source; unknown markup remains visible.
    parts=re.split(r'(`[^`]+`|\[[^\]]+\]\(https?://[^)]+\))',text)
    for part in parts:
        if not part: continue
        if part.startswith('`'):
            run=p.add_run(part[1:-1]); run.font.name='Liberation Mono'; run.font.size=Pt(9)
        elif re.fullmatch(r'\[[^\]]+\]\(https?://[^)]+\)',part):
            label,url=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',part).groups()
            h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
            r=OxmlElement('w:r'); rp=OxmlElement('w:rPr'); c=OxmlElement('w:color');c.set(qn('w:val'),'1F4E79');rp.append(c)
            r.append(rp);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)
        else: p.add_run(part)


def make_table(doc, rows):
    n=len(rows[0]); table=doc.add_table(rows=1, cols=n)
    table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    if n==7: widths=[.57,.65,.62,1.18,1.18,1.18,1.22]
    elif n==4: widths=[2.05,1.95,1.4,1.2]
    elif rows[0][0]=='编号': widths=[.55,3.85,2.2]
    elif rows[0][0]=='分支': widths=[1.1,3.0,2.5]
    elif rows[0][0]=='环节': widths=[1.0,2.6,3.0]
    elif rows[0][0]=='数值诊断指标': widths=[3.6,1.5,1.5]
    elif rows[0][0]=='规定候选动作': widths=[3.15,1.1,2.35]
    else: widths=[1.35,1.3,3.95]
    widths=[w*6.9/sum(widths) for w in widths]
    for col,w in zip(table.columns,widths): col.width=Inches(w)
    pr=table._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+edge)
        for k,v in [('val','single'),('sz','4'),('color','D9D9D9')]:e.set(qn('w:'+k),v)
        borders.append(e)
    pr.append(borders)
    margin=OxmlElement('w:tblCellMar')
    for edge in ('top','left','bottom','right'):
        e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');margin.append(e)
    pr.append(margin)
    for ri,values in enumerate(rows):
        row=table.rows[0] if ri==0 else table.add_row()
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if ri==0: trpr.append(OxmlElement('w:tblHeader'))
        for ci,(cell,value) in enumerate(zip(row.cells,values)):
            cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.style=doc.styles['Table Text']
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER if n==7 or (ci==1 and rows[0][0] in ('校准预算','规定候选动作')) or (ci>=1 and rows[0][0]=='数值诊断指标') else WD_ALIGN_PARAGRAPH.LEFT
            inline(p,value)
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EDF3');cell._tc.get_or_add_tcPr().append(shade)
                for run in p.runs:run.bold=True
    doc.add_paragraph().paragraph_format.space_after=Pt(2)


def build():
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.78)
    sec.left_margin=sec.right_margin=Inches(.8)
    sec.footer_distance=Inches(.35)
    set_font(doc.styles['Normal'],11)
    pf=doc.styles['Normal'].paragraph_format
    pf.line_spacing=1.25;pf.space_after=Pt(7);pf.widow_control=True
    set_font(doc.styles['Title'],23,'Noto Sans CJK SC')
    doc.styles['Title'].paragraph_format.space_after=Pt(12)
    set_font(doc.styles['Subtitle'],13,'Noto Sans CJK SC')
    doc.styles['Subtitle'].font.italic=False
    # The bundled default document carries a Title border; remove it explicitly.
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    for style,size in [('Heading 1',15),('Heading 2',12)]:
        set_font(doc.styles[style],size,'Noto Sans CJK SC')
        doc.styles[style].font.bold=True
        f=doc.styles[style].paragraph_format;f.space_before=Pt(14);f.space_after=Pt(7);f.keep_with_next=True
    ts=doc.styles.add_style('Table Text',1);set_font(ts,9.5)
    ts.paragraph_format.line_spacing=1.16;ts.paragraph_format.space_after=Pt(2);ts.paragraph_format.space_before=Pt(2)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=footer.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');run._r.addnext(fld)
    lines=(HERE/'CPR2_E0D_TECHNICAL_REPORT.md').read_text().splitlines()
    i=0;eq=0
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line=='$$':
            i+=1
            while lines[i].strip()!='$$':i+=1
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p._p.append(node('oMathPara',[node('oMath',equation(eq))]));eq+=1
        elif line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',c) for c in cells): rows.append(cells)
                i+=1
            make_table(doc,rows);continue
        elif line.startswith('# '):doc.add_paragraph(line[2:],style='Title')
        elif line=='## E0 D 阶段技术报告':doc.add_paragraph(line[3:],style='Subtitle')
        elif line.startswith('### '):doc.add_paragraph(line[4:],style='Heading 2')
        elif line.startswith('## '):
            p=doc.add_paragraph(line[3:],style='Heading 1')
            if line.startswith('## 附录 A'):p.paragraph_format.page_break_before=True
        else:
            p=doc.add_paragraph();inline(p,line)
        i+=1
    assert eq==4
    doc.core_properties.title='CPR 2 点云修订方法的分解实验与研究价值评估'
    doc.core_properties.subject='E0 D 阶段技术报告'
    doc.core_properties.author='点云修订研究项目'
    doc.core_properties.keywords='CPR-2, E0-D, calibration, projection, point cloud'
    doc.save(HERE/'CPR2_E0D_TECHNICAL_REPORT.docx')
    print('Built DOCX with',len(doc.paragraphs),'paragraphs,',len(doc.tables),'tables,',eq,'native math blocks')


if __name__=='__main__':
    recheck();build()
